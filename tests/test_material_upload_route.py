import io
import os
import sqlite3
import tempfile
import unittest
from unittest.mock import MagicMock, patch

import pypdf
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas

# app.py initializes the configured database at import time. Suppress that
# initialization when this test imports the app so this test never touches it.
import database

with patch.object(database, 'init_db'):
    import app as app_module

from app import app
from services.embedding_service import EmbeddingService
from tests.test_embeddings import MockEmbedder


class TestMaterialUploadRoute(unittest.TestCase):
    """Exercise multipart PDF upload through extraction, storage, indexing, and search."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.upload_dir = os.path.join(self.temp_dir.name, 'uploads')
        self.database_path = os.path.join(self.temp_dir.name, 'upload_test.sqlite')
        self.vector_dir = os.path.join(self.temp_dir.name, 'vector_store')
        self.index_path = os.path.join(self.vector_dir, 'chunks.faiss')
        self.metadata_path = os.path.join(self.vector_dir, 'chunks_metadata.json')
        self._initialize_test_database()

        self.embedding_service = EmbeddingService(
            vector_dir=self.vector_dir,
            index_path=self.index_path,
            metadata_path=self.metadata_path,
            model=MockEmbedder(dimension=384),
            project_root=self.temp_dir.name,
        )

        self.material_service = app_module.material_service
        self.original_upload_dir = self.material_service.upload_folder
        self.original_database_path = self.material_service.database_path
        self.original_material_embedding_service = self.material_service._embedding_service
        self.original_shared_embedding_service = app_module.embedding_service
        self.original_max_content_length = app.config['MAX_CONTENT_LENGTH']

        self.material_service.upload_folder = self.upload_dir
        self.material_service.database_path = self.database_path
        self.material_service._embedding_service = self.embedding_service
        app_module.embedding_service = self.embedding_service

        app.config['TESTING'] = True
        self.client = app.test_client()

    def tearDown(self):
        self.material_service.upload_folder = self.original_upload_dir
        self.material_service.database_path = self.original_database_path
        self.material_service._embedding_service = self.original_material_embedding_service
        app_module.embedding_service = self.original_shared_embedding_service
        app.config['MAX_CONTENT_LENGTH'] = self.original_max_content_length
        self.temp_dir.cleanup()

    def _initialize_test_database(self):
        connection = sqlite3.connect(self.database_path)
        connection.executescript(
            '''
            CREATE TABLE users (
                id INTEGER PRIMARY KEY,
                name TEXT NOT NULL,
                email TEXT UNIQUE NOT NULL,
                password_hash TEXT NOT NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            );
            CREATE TABLE materials (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                filename TEXT NOT NULL,
                file_type TEXT NOT NULL,
                file_path TEXT NOT NULL,
                file_size INTEGER,
                page_count INTEGER,
                processing_status TEXT DEFAULT 'pending',
                uploaded_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            );
            CREATE TABLE document_chunks (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                material_id INTEGER NOT NULL,
                chunk_index INTEGER NOT NULL,
                content TEXT NOT NULL,
                source_type TEXT NOT NULL,
                source_file TEXT NOT NULL,
                page_number INTEGER NOT NULL,
                location_type TEXT NOT NULL DEFAULT 'page',
                location_value TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            );
            INSERT INTO users (id, name, email, password_hash)
            VALUES (1, 'Test User', 'upload-test@example.com', 'test-hash');
            '''
        )
        connection.commit()
        connection.close()

    @staticmethod
    def _create_two_page_pdf(
        page_one='Page one topic: photosynthesis converts sunlight into chemical energy.',
        page_two='Page two topic: mitochondria produce cellular energy through respiration.',
    ):
        pdf_buffer = io.BytesIO()
        document = canvas.Canvas(pdf_buffer)
        document.drawString(72, 720, page_one)
        document.showPage()
        document.drawString(72, 720, page_two)
        document.showPage()
        document.save()
        pdf_buffer.seek(0)
        return pdf_buffer

    @staticmethod
    def _create_sample_pptx():
        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[6])

        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        textbox.text_frame.text = 'Slide two: photosynthesis uses sunlight.'

        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        table = slide.shapes.add_table(
            1, 2, Inches(1), Inches(2), Inches(5), Inches(1)
        ).table
        table.cell(0, 0).text = 'Slide three:'
        table.cell(0, 1).text = 'chlorophyll captures light energy.'

        pptx_buffer = io.BytesIO()
        presentation.save(pptx_buffer)
        pptx_buffer.seek(0)
        return pptx_buffer

    def test_upload_size_limits_are_aligned_at_50_mb(self):
        expected_limit = 50 * 1024 * 1024

        self.assertEqual(app.config['MAX_CONTENT_LENGTH'], expected_limit)
        self.assertEqual(self.material_service.MAX_FILE_SIZE, expected_limit)

    def test_material_page_displays_configured_limit_and_accessible_progress_status(self):
        response = self.client.get('/materials')

        self.assertEqual(response.status_code, 200)
        self.assertIn(b'PDF, PPTX, and timestamped TXT transcripts (Max 50 MB)', response.data)
        self.assertIn(b'[HH:MM:SS] spoken text', response.data)
        self.assertIn(b'Direct video transcription is not supported', response.data)
        self.assertNotIn(b'Video (Coming Soon)', response.data)
        self.assertIn(b'accept=".pdf,.pptx,.txt"', response.data)
        self.assertNotIn(b'16MB', response.data)
        self.assertIn(b'id="uploadProgress"', response.data)
        self.assertIn(b'role="status"', response.data)
        self.assertIn(b'aria-live="polite"', response.data)

    def test_material_service_rejects_file_above_configured_limit(self):
        oversized_file = MagicMock()
        oversized_file.filename = 'oversized.pdf'
        oversized_file.tell.return_value = app.config['MAX_CONTENT_LENGTH'] + 1

        file_path, error = self.material_service.save_uploaded_file(oversized_file, user_id=1)

        self.assertIsNone(file_path)
        self.assertEqual(error, 'File too large. Maximum size is 50MB.')
        self.assertFalse(os.path.exists(self.upload_dir) and os.listdir(self.upload_dir))

    def test_material_service_rejects_oversized_pptx(self):
        oversized_file = MagicMock()
        oversized_file.filename = 'oversized.PPTX'
        oversized_file.tell.return_value = app.config['MAX_CONTENT_LENGTH'] + 1

        file_path, error = self.material_service.save_uploaded_file(oversized_file, user_id=1)

        self.assertIsNone(file_path)
        self.assertEqual(error, 'File too large. Maximum size is 50MB.')
        self.assertFalse(os.path.exists(self.upload_dir) and os.listdir(self.upload_dir))

    def test_unsupported_upload_is_rejected_without_saving_or_creating_record(self):
        response = self.client.post(
            '/materials',
            data={'file': (io.BytesIO(b'not supported'), 'notes.csv')},
            content_type='multipart/form-data',
            follow_redirects=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn(b'Only PDF, PowerPoint (.pptx), and timestamped transcript (.txt) files are allowed', response.data)
        self.assertFalse(os.path.exists(self.upload_dir) and os.listdir(self.upload_dir))
        connection = sqlite3.connect(self.database_path)
        material_count = connection.execute('SELECT COUNT(*) FROM materials').fetchone()[0]
        connection.close()
        self.assertEqual(material_count, 0)

    def test_multipart_pptx_upload_stores_slide_locations_and_incrementally_indexes(self):
        pptx = self._create_sample_pptx()
        index_material = self.embedding_service.index_material
        with patch.object(
            self.embedding_service, 'index_material', wraps=index_material
        ) as index_material_spy:
            response = self.client.post(
                '/materials',
                data={'file': (pptx, 'course_slides.PPTX')},
                content_type='multipart/form-data',
                follow_redirects=True,
            )

        self.assertEqual(response.status_code, 200)
        self.assertIn(b'Successfully uploaded and processed course_slides.PPTX', response.data)
        self.assertEqual(index_material_spy.call_count, 1)
        uploaded_path = os.path.join(self.upload_dir, 'course_slides.PPTX')
        self.assertTrue(os.path.isfile(uploaded_path))

        connection = sqlite3.connect(self.database_path)
        connection.row_factory = sqlite3.Row
        material = connection.execute(
            'SELECT id, filename, file_type, page_count, processing_status FROM materials'
        ).fetchone()
        chunks = connection.execute(
            '''SELECT id, chunk_index, source_type, source_file, page_number,
                      location_type, location_value, content
               FROM document_chunks WHERE material_id = ? ORDER BY chunk_index''',
            (material['id'],),
        ).fetchall()
        connection.close()

        self.assertEqual(material['filename'], 'course_slides.PPTX')
        self.assertEqual(material['file_type'], 'pptx')
        self.assertEqual(material['page_count'], 3)
        self.assertEqual(material['processing_status'], 'processed')
        self.assertEqual([row['page_number'] for row in chunks], [2, 3])
        self.assertEqual([row['location_type'] for row in chunks], ['slide', 'slide'])
        self.assertEqual([row['location_value'] for row in chunks], ['2', '3'])
        self.assertEqual([row['source_type'] for row in chunks], ['slides', 'slides'])
        self.assertTrue(all(row['source_file'] == 'course_slides.PPTX' for row in chunks))
        self.assertIn('photosynthesis', chunks[0]['content'])
        self.assertIn('chlorophyll', chunks[1]['content'])
        self.assertEqual(self.embedding_service.index.ntotal, len(chunks))
        self.assertEqual(
            {metadata['location_value'] for metadata in self.embedding_service.metadata},
            {'2', '3'},
        )
        self.assertTrue(all(
            metadata['location_type'] == 'slide'
            for metadata in self.embedding_service.metadata
        ))
        retrieved = self.embedding_service.search(
            'photosynthesis sunlight', top_k=5, material_id=material['id'], min_score=0.0
        )
        self.assertTrue(retrieved)
        self.assertTrue(all(item['location_type'] == 'slide' for item in retrieved))

    def test_multipart_transcript_upload_stores_timestamps_and_incrementally_indexes(self):
        transcript = io.BytesIO(
            b'[00:00:12] Introduction to database systems.\n'
            b'[00:01:25] A DBMS organizes, stores, and retrieves data.\n'
        )
        index_material = self.embedding_service.index_material
        with patch.object(
            self.embedding_service, 'index_material', wraps=index_material
        ) as index_material_spy:
            response = self.client.post(
                '/materials',
                data={'file': (transcript, 'lecture_transcript.TXT')},
                content_type='multipart/form-data',
                follow_redirects=True,
            )

        self.assertEqual(response.status_code, 200)
        self.assertIn(b'Successfully uploaded and processed lecture_transcript.TXT', response.data)
        self.assertEqual(index_material_spy.call_count, 1)
        self.assertTrue(os.path.isfile(os.path.join(self.upload_dir, 'lecture_transcript.TXT')))

        connection = sqlite3.connect(self.database_path)
        connection.row_factory = sqlite3.Row
        material = connection.execute(
            'SELECT id, file_type, page_count, processing_status FROM materials'
        ).fetchone()
        chunks = connection.execute(
            '''SELECT page_number, location_type, location_value, source_type
               FROM document_chunks WHERE material_id = ? ORDER BY chunk_index''',
            (material['id'],),
        ).fetchall()
        connection.close()

        self.assertEqual(material['file_type'], 'transcript')
        self.assertEqual(material['page_count'], 2)
        self.assertEqual(material['processing_status'], 'processed')
        self.assertEqual([row['page_number'] for row in chunks], [1, 1])
        self.assertEqual([row['location_type'] for row in chunks], ['timestamp', 'timestamp'])
        self.assertEqual([row['location_value'] for row in chunks], ['00:00:12', '00:01:25'])
        self.assertEqual([row['source_type'] for row in chunks], ['transcript', 'transcript'])
        self.assertEqual(self.embedding_service.index.ntotal, 2)
        self.assertEqual(
            {metadata['location_value'] for metadata in self.embedding_service.metadata},
            {'00:00:12', '00:01:25'},
        )
        self.assertTrue(all(
            metadata['location_type'] == 'timestamp'
            for metadata in self.embedding_service.metadata
        ))
        retrieved = self.embedding_service.search(
            'DBMS stores retrieves data', top_k=2, material_id=material['id'], min_score=0.0
        )
        self.assertTrue(retrieved)
        self.assertTrue(all(result['location_type'] == 'timestamp' for result in retrieved))

    def test_empty_transcript_upload_is_rejected_with_a_clear_message(self):
        response = self.client.post(
            '/materials',
            data={'file': (io.BytesIO(b' \n\n'), 'empty_transcript.txt')},
            content_type='multipart/form-data',
            follow_redirects=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn(b'Transcript contains no timestamped text', response.data)
        connection = sqlite3.connect(self.database_path)
        material_count = connection.execute('SELECT COUNT(*) FROM materials').fetchone()[0]
        connection.close()
        self.assertEqual(material_count, 0)

    def test_invalid_transcript_upload_is_not_marked_processed(self):
        response = self.client.post(
            '/materials',
            data={'file': (io.BytesIO(b'not a timestamped transcript'), 'invalid.txt')},
            content_type='multipart/form-data',
            follow_redirects=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn(b'Processing failed', response.data)
        connection = sqlite3.connect(self.database_path)
        material_count = connection.execute('SELECT COUNT(*) FROM materials').fetchone()[0]
        connection.close()
        self.assertEqual(material_count, 0)

    def test_material_service_rejects_oversized_transcript(self):
        oversized_file = MagicMock()
        oversized_file.filename = 'oversized.TXT'
        oversized_file.tell.return_value = app.config['MAX_CONTENT_LENGTH'] + 1

        file_path, error = self.material_service.save_uploaded_file(oversized_file, user_id=1)

        self.assertIsNone(file_path)
        self.assertEqual(error, 'File too large. Maximum size is 50MB.')

    def test_pptx_indexing_failure_does_not_mark_material_processed(self):
        failing_index = MagicMock()
        failing_index.index_material.side_effect = RuntimeError('simulated indexing failure')
        self.material_service._embedding_service = failing_index

        response = self.client.post(
            '/materials',
            data={'file': (self._create_sample_pptx(), 'failure_slides.pptx')},
            content_type='multipart/form-data',
            follow_redirects=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn(b'semantic indexing failed', response.data)
        connection = sqlite3.connect(self.database_path)
        material = connection.execute(
            'SELECT id, processing_status FROM materials'
        ).fetchone()
        chunk_count = connection.execute(
            'SELECT COUNT(*) FROM document_chunks WHERE material_id = ?',
            (material[0],),
        ).fetchone()[0]
        connection.close()
        self.assertEqual(material[1], 'indexing_failed')
        self.assertGreater(chunk_count, 0)
        self.assertTrue(os.path.isfile(os.path.join(self.upload_dir, 'failure_slides.pptx')))

    def test_upload_save_exception_returns_user_friendly_flash_message(self):
        with patch.object(
            self.material_service,
            'save_uploaded_file',
            side_effect=OSError('private filesystem detail'),
        ):
            response = self.client.post(
                '/materials',
                data={'file': (io.BytesIO(b'pdf data'), 'notes.pdf')},
                content_type='multipart/form-data',
                follow_redirects=True,
            )

        self.assertEqual(response.status_code, 200)
        self.assertIn(b'We could not save the material.', response.data)
        self.assertNotIn(b'private filesystem detail', response.data)

    def test_flask_rejects_oversized_upload_with_user_friendly_message(self):
        app.config['MAX_CONTENT_LENGTH'] = 2 * 1024 * 1024

        response = self.client.post(
            '/materials',
            data={'file': (io.BytesIO(b'x' * (2 * 1024 * 1024 + 1)), 'oversized.pdf')},
            content_type='multipart/form-data',
            follow_redirects=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn(b'The uploaded file is too large.', response.data)
        self.assertIn(b'maximum upload size is 2 MB', response.data)
        self.assertFalse(os.path.exists(self.upload_dir) and os.listdir(self.upload_dir))
        connection = sqlite3.connect(self.database_path)
        material_count = connection.execute('SELECT COUNT(*) FROM materials').fetchone()[0]
        connection.close()
        self.assertEqual(material_count, 0)

    def test_multipart_upload_extracts_stores_indexes_and_retrieves_both_pages(self):
        pdf = self._create_two_page_pdf()
        original_pdf_reader = pypdf.PdfReader
        with patch(
            'ingestion.pdf_processor.pypdf.PdfReader',
            wraps=original_pdf_reader,
        ) as reader_factory:
            with self.assertLogs(app.logger.name, level='INFO') as captured_logs:
                response = self.client.post(
                    '/materials',
                    data={'file': (pdf, 'two_page_notes.pdf')},
                    content_type='multipart/form-data',
                    follow_redirects=True,
                )

        self.assertEqual(reader_factory.call_count, 1)

        self.assertEqual(response.status_code, 200)
        self.assertIn(b'Successfully uploaded and processed two_page_notes.pdf', response.data)
        log_output = '\n'.join(captured_logs.output)
        self.assertIn('PDF text extraction completed', log_output)
        self.assertIn('SQLite chunk storage completed', log_output)
        self.assertIn('FAISS indexing completed', log_output)
        self.assertNotIn('photosynthesis converts', log_output)

        uploaded_path = os.path.realpath(os.path.join(self.upload_dir, 'two_page_notes.pdf'))
        self.assertTrue(os.path.isfile(uploaded_path))
        self.assertTrue(uploaded_path.startswith(os.path.realpath(self.upload_dir) + os.sep))
        self.assertTrue(os.path.realpath(self.database_path).startswith(
            os.path.realpath(self.temp_dir.name) + os.sep
        ))
        self.assertTrue(os.path.realpath(self.index_path).startswith(
            os.path.realpath(self.temp_dir.name) + os.sep
        ))

        connection = sqlite3.connect(self.database_path)
        connection.row_factory = sqlite3.Row
        material = connection.execute(
            'SELECT id, filename, page_count, processing_status FROM materials'
        ).fetchone()
        chunks = connection.execute(
            '''SELECT id, chunk_index, source_file, page_number, content
               FROM document_chunks WHERE material_id = ? ORDER BY page_number''',
            (material['id'],),
        ).fetchall()
        connection.close()

        self.assertEqual(material['filename'], 'two_page_notes.pdf')
        self.assertEqual(material['page_count'], 2)
        self.assertEqual(material['processing_status'], 'processed')
        self.assertEqual([row['page_number'] for row in chunks], [1, 2])
        self.assertIn('photosynthesis', chunks[0]['content'])
        self.assertIn('mitochondria', chunks[1]['content'])
        self.assertEqual(self.embedding_service.index.ntotal, 2)
        self.assertEqual(len(self.embedding_service.metadata), 2)

        metadata_by_chunk_id = {
            item['chunk_id']: item for item in self.embedding_service.metadata
        }
        self.assertEqual(
            {metadata_by_chunk_id[row['id']]['page_number'] for row in chunks},
            {1, 2},
        )
        self.assertTrue(all(
            metadata_by_chunk_id[row['id']]['source_file'] == 'two_page_notes.pdf'
            for row in chunks
        ))

        search_results = self.embedding_service.search(
            'photosynthesis sunlight chemical energy',
            top_k=5,
            material_id=material['id'],
            min_score=0.0,
        )
        self.assertEqual({result['page_number'] for result in search_results}, {1, 2})
        self.assertTrue(all(
            result['source_file'] == 'two_page_notes.pdf'
            and result['material_id'] == material['id']
            for result in search_results
        ))
        self.assertEqual(
            {result['page_number']: result['chunk_id'] for result in search_results},
            {row['page_number']: row['id'] for row in chunks},
        )

    def test_duplicate_filename_uploads_preserve_both_materials_and_files(self):
        filename = 'duplicate_notes.pdf'
        first_pdf = self._create_two_page_pdf(
            'First upload page one: photosynthesis in green plants.',
            'First upload page two: chlorophyll absorbs light.',
        )
        first_response = self.client.post(
            '/materials',
            data={'file': (first_pdf, filename)},
            content_type='multipart/form-data',
            follow_redirects=True,
        )
        self.assertEqual(first_response.status_code, 200)
        self.assertIn(b'Successfully uploaded and processed duplicate_notes.pdf', first_response.data)

        first_path = os.path.join(self.upload_dir, filename)
        with open(first_path, 'rb') as saved_first:
            first_bytes = saved_first.read()

        second_pdf = self._create_two_page_pdf(
            'Second upload page one: cellular respiration.',
            'Second upload page two: mitochondria produce ATP.',
        )
        second_response = self.client.post(
            '/materials',
            data={'file': (second_pdf, filename)},
            content_type='multipart/form-data',
            follow_redirects=True,
        )
        self.assertEqual(second_response.status_code, 200)
        self.assertIn(b'Successfully uploaded and processed duplicate_notes.pdf', second_response.data)

        second_path = os.path.join(self.upload_dir, 'duplicate_notes_1.pdf')
        self.assertTrue(os.path.isfile(first_path))
        self.assertTrue(os.path.isfile(second_path))
        with open(first_path, 'rb') as saved_first:
            self.assertEqual(saved_first.read(), first_bytes)
        with open(second_path, 'rb') as saved_second:
            second_bytes = saved_second.read()
        self.assertNotEqual(first_bytes, second_bytes)

        connection = sqlite3.connect(self.database_path)
        connection.row_factory = sqlite3.Row
        materials = connection.execute(
            'SELECT id, filename, file_path, processing_status FROM materials ORDER BY id'
        ).fetchall()
        connection.close()

        self.assertEqual(len(materials), 2)
        self.assertEqual([row['filename'] for row in materials], [filename, filename])
        self.assertEqual([row['processing_status'] for row in materials], ['processed', 'processed'])
        self.assertNotEqual(materials[0]['file_path'], materials[1]['file_path'])
        self.assertEqual(os.path.realpath(materials[0]['file_path']), os.path.realpath(first_path))
        self.assertEqual(os.path.realpath(materials[1]['file_path']), os.path.realpath(second_path))

        self.assertEqual(self.embedding_service.index.ntotal, 4)
        for material in materials:
            results = self.embedding_service.search(
                'cellular energy',
                top_k=5,
                material_id=material['id'],
                min_score=0.0,
            )
            self.assertEqual({item['page_number'] for item in results}, {1, 2})
            self.assertTrue(all(item['material_id'] == material['id'] for item in results))

    def test_embedding_service_initialization_is_timed_without_loading_model(self):
        original_shared_service = app_module.embedding_service
        app_module.embedding_service = None
        try:
            with patch.object(app_module, 'EmbeddingService', return_value=self.embedding_service):
                with self.assertLogs(app.logger.name, level='INFO') as captured_logs:
                    initialized_service = app_module.get_embedding_service()
        finally:
            app_module.embedding_service = original_shared_service

        self.assertIs(initialized_service, self.embedding_service)
        self.assertIn('Embedding service initialization started', '\n'.join(captured_logs.output))
        self.assertIn('Embedding service initialization completed', '\n'.join(captured_logs.output))


if __name__ == '__main__':
    unittest.main()
